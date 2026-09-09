import tkinter as tk
from tkinter import messagebox
from selenium import webdriver
from selenium.webdriver.chrome.service import Service as ChromeService
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager
from webdriver_manager.core.os_manager import ChromeType
import time
import re

def resolve_chromedriver():
    """Reusa somente um driver já instalado e compatível com este Chromium."""
    import json
    import os
    from pathlib import Path
    import subprocess
    import tempfile
    def version(path):
        text = subprocess.check_output([str(path), '--version'], text=True, timeout=5)
        found = re.search(r'\b(\d+\.\d+\.\d+\.\d+)\b', text)
        if not found:raise ValueError('Versão do executável não reconhecida.')
        return found[1]
    cache = Path(__file__).resolve().parents[1] / 'data' / 'chromedriver-booking.json'
    try:browser_version = version('/usr/bin/chromium')
    except Exception:browser_version = None
    if browser_version:
        try:
            saved = json.loads(cache.read_text())
            path = Path(saved['path'])
            allowed_root = (Path.home() / '.wdm').resolve()
            if (saved['browser_version'] == browser_version and path.is_absolute()
                    and path.resolve().is_relative_to(allowed_root) and path.is_file()
                    and version(path).split('.')[:3] == browser_version.split('.')[:3]):
                print('ChromeDriver local conferido; reutilizando instalação.', flush=True)
                return str(path)
        except Exception:
            pass
    path = ChromeDriverManager(chrome_type=ChromeType.CHROMIUM).install()
    if browser_version:
        try:
            if version(path).split('.')[:3] == browser_version.split('.')[:3]:
                cache.parent.mkdir(exist_ok=True)
                with tempfile.NamedTemporaryFile(mode='w',dir=cache.parent,delete=False) as f:
                    tmp = Path(f.name)
                    json.dump({'browser_version':browser_version,'path':str(Path(path).resolve())},f)
                try:os.replace(tmp,cache)
                finally:tmp.unlink(missing_ok=True)
        except Exception:
            pass
    return path


def create_driver(headless=False):
    options = webdriver.ChromeOptions()

    # Chromium no Arch
    options.binary_location = "/usr/bin/chromium"

    if headless:
        options.add_argument("--headless=new")
    options.page_load_strategy = "eager"

    prefs = {"profile.managed_default_content_settings.images": 2}
    options.add_experimental_option("prefs", prefs)

    # Flags importantes no Linux / Arch
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-blink-features=AutomationControlled")
    options.add_experimental_option("excludeSwitches", ["enable-automation"])
    options.add_argument("--window-size=1920,1080")

    service = ChromeService(resolve_chromedriver())
    return webdriver.Chrome(service=service, options=options)

def scroll_until_all_hotels_loaded(driver, max_wait_time=60, total_timeout=None):
    SCROLL_PAUSE_TIME = 2.5
    last_count = 0
    start_time = time.monotonic()
    began = start_time

    while True:
        driver.execute_script("window.scrollBy(0, 2000);")
        time.sleep(SCROLL_PAUSE_TIME)

        hotels = driver.find_elements(By.XPATH, '//div[@data-testid="property-card"]')
        current_count = len(hotels)
        print(f"Hotéis visíveis: {current_count}")

        if current_count > last_count:
            last_count = current_count
            start_time = time.monotonic()
        else:
            if time.monotonic() - start_time > max_wait_time:
                print(f"⏹️ Parou de rolar: nenhum hotel novo após {max_wait_time} segundos.", flush=True)
                break

        if total_timeout is not None and time.monotonic() - began >= total_timeout:
            print("Janela de coleta atingida; cobertura parcial dos cartões carregados.", flush=True)
            break

    return hotels

NOT_INFORMED = "Não informado"
REFERENCE_COORDS = (-22.5113, -43.1779)
REFERENCE_NAME = "Ponto de referência em Petrópolis (-22.5113, -43.1779)"


def first_text(root, selectors):
    for selector in selectors:
        for element in root.find_elements(By.CSS_SELECTOR, selector):
            text = element.text.strip()
            if text:
                return text
    return ""


def parse_price(raw):
    """Aceita somente valor identificado como BRL; mantém o texto para auditoria."""
    match = re.fullmatch(r"\s*(?:R\$|BRL)\s*([0-9]+(?:\.[0-9]{3})*(?:,[0-9]{2})?)\s*", raw)
    if not match:
        return None, NOT_INFORMED
    value = float(match.group(1).replace(".", "").replace(",", "."))
    return (value, "BRL") if value > 0 else (None, "BRL")


def price_with_taxes(price, currency, taxes_text):
    """Total dos valores exibidos; não presume taxa zero quando não informada."""
    if price is None or currency != 'BRL':
        return None, None, 'Preço ou moeda não confirmado'
    if re.fullmatch(r'\s*Impostos e taxas incluídos\s*', taxes_text, re.I):
        return 0.0, price, 'Taxas incluídas no preço exibido'
    match = re.fullmatch(
        r'\s*\+\s*(R\$\s*[0-9]+(?:\.[0-9]{3})*(?:,[0-9]{2})?)\s+em impostos e taxas\s*',
        taxes_text, re.I)
    if match:
        extra, _ = parse_price(match.group(1))
        if extra is not None:
            return extra, round(price + extra, 2), 'Preço exibido + taxas adicionais exibidas'
    return None, None, 'Total não confirmado: verificar taxas na página'


def extract_stars(card):
    for container in card.find_elements(By.CSS_SELECTOR, '[data-testid="rating-stars"]'):
        for element in [container] + container.find_elements(By.CSS_SELECTOR, '[aria-label], [title]'):
            for attribute in ('aria-label', 'title'):
                label = element.get_attribute(attribute) or ""
                match = re.search(r'\b([1-5])\s*(?:estrelas?|stars?)\b', label, re.I)
                if match:
                    return int(match.group(1))
    return NOT_INFORMED


def read_cards_batch(driver):
    """Uma chamada ao navegador para ler o lote inteiro."""
    raw = driver.execute_script('''
      const rows = [...document.querySelectorAll('[data-testid="property-card"]')].map(card => {
        const text = selector => (card.querySelector(selector)?.innerText || '').trim();
        return {
          'Hotel Name': text('[data-testid="title"]'),
          'Hotel URL': card.querySelector('[data-testid="title-link"]')?.href || '',
          'Price Text': text('[data-testid="price-and-discounted-price"]'),
          'Taxes Text': text('[data-testid="taxes-and-charges"]'),
          'Stay Text': text('[data-testid="price-for-x-nights"]'),
          'Room Text': text('[data-testid="recommended-unit"]'),
          'Review Text': text('[data-testid="review-score"]')
        };
      });
      return {url: location.href, rows};
    ''')
    rows = []
    for row in raw['rows']:
        price, currency = parse_price(row.get('Price Text', ''))
        row.update({'Price': price, 'Currency': currency, 'Stars': NOT_INFORMED,
                    'Review Score (/10)': None,
                    'Price Status': 'Extraído do cartão' if price is not None else 'Não confirmado'})
        match = re.search(r'(?:Com nota|Pontuação|Scored|Nota)\s*([0-9]+(?:[.,][0-9])?)', row.pop('Review Text', ''), re.I)
        if match and 0 <= float(match[1].replace(',', '.')) <= 10:
            row['Review Score (/10)'] = float(match[1].replace(',', '.'))
        rows.append(row)
    return rows, raw['url']


def stable_batch(rows, seen, now, stable_seconds=3):
    from urllib.parse import urlsplit
    current, stable = {}, []
    for row in rows:
        u = urlsplit(row.get('Hotel URL', ''))
        if u.scheme != 'https' or u.hostname not in ('www.booking.com', 'booking.com') or not row.get('Hotel Name'):
            continue
        key = u.hostname + u.path
        signature = tuple(row.get(k) for k in ('Hotel Name', 'Price Text', 'Taxes Text', 'Stay Text', 'Room Text'))
        previous = seen.get(key)
        since = previous[1] if previous and previous[0] == signature else now
        current[key] = (signature, since)
        if row.get('Price') is not None and now - since >= stable_seconds:
            stable.append(row)
    seen.clear(); seen.update(current)
    return stable


def extract_fast_batch(driver, on_progress=None, confirm=None):
    import json
    began = time.monotonic()
    seen, first, last_signature = {}, None, None
    count, last_growth = 0, began
    last_scroll = began
    while True:
        rows, current_url = read_cards_batch(driver)
        if confirm and not confirm(current_url):
            raise RuntimeError('Consulta mudou durante a coleta; ofertas não confirmadas.')
        now = time.monotonic()
        if len(rows) > count:
            count, last_growth = len(rows), now
            print('Hotéis carregados: ' + str(count), flush=True)
        stable = stable_batch(rows, seen, now)
        signature = json.dumps(stable, sort_keys=True, ensure_ascii=False)
        if signature != last_signature and (stable or first is not None):
            if stable and first is None:
                first = now - began
                print('Primeiro lote conferido em ' + str(round(first, 2)) + ' s após iniciar leitura dos cartões.', flush=True)
            if on_progress: on_progress(stable)
            last_signature = signature
        # Primeiro lote antes da rolagem; depois carrega o restante sem visitas individuais.
        if first is not None and now - last_scroll >= 2.5:
            driver.execute_script('window.scrollBy(0, 2000)')
            last_scroll = now
        if now - began >= 45 or (first is not None and now - last_growth >= 12):
            if not stable:
                raise RuntimeError('Nenhum preço estável disponível ao encerrar a coleta.')
            print('Coleta encerrada com ' + str(len(stable)) + ' ofertas estáveis; cobertura parcial.', flush=True)
            return stable
        time.sleep(.5)


def extract_hotels(driver, fast=False, on_progress=None, confirm=None):
    if fast:
        return extract_fast_batch(driver, on_progress=on_progress, confirm=confirm)
    hotels = scroll_until_all_hotels_loaded(driver, max_wait_time=12 if fast else 60, total_timeout=45 if fast else None)
    hotel_list = []
    for hotel in hotels:
        raw_price = first_text(hotel, ['[data-testid="price-and-discounted-price"]'])
        price, currency = parse_price(raw_price)
        score = None
        # Preserva o seletor que funcionou no primeiro teste; valida a faixa.
        raw_score = first_text(hotel, ['div.f63b14ab7a'])
        if re.fullmatch(r'\d{1,2}(?:[.,]\d)?', raw_score):
            score = float(raw_score.replace(',', '.'))
        if score is None or not 0 <= score <= 10:
            raw_score = first_text(hotel, ['[data-testid="review-score"]'])
            match = re.search(r'(?:Com nota|Pontuação|Scored|Nota)\s*([0-9]+(?:[.,][0-9])?)', raw_score, re.I)
            score = float(match.group(1).replace(',', '.')) if match else None
        if score is not None and not 0 <= score <= 10:
            score = None
        links = hotel.find_elements(By.CSS_SELECTOR, '[data-testid="title-link"]')
        url = links[0].get_attribute('href') if links else None
        hotel_list.append({
            'Hotel Name': first_text(hotel, ['[data-testid="title"]']) or NOT_INFORMED,
            'Stars': extract_stars(hotel),
            'Price': price,
            'Currency': currency,
            'Price Text': raw_price or NOT_INFORMED,
            'Price Status': 'Extraído do cartão' if price is not None else 'Não confirmado',
            'Stay Text': first_text(hotel, ['[data-testid="price-for-x-nights"]']) or NOT_INFORMED,
            'Taxes Text': first_text(hotel, ['[data-testid="taxes-and-charges"]']) or NOT_INFORMED,
            'Room Text': first_text(hotel, ['[data-testid="recommended-unit"]']) or NOT_INFORMED,
            'Review Score (/10)': score,
            'Hotel URL': url or NOT_INFORMED,
        })
    return hotel_list


def address_from_jsonld(value):
    """Lê PostalAddress de uma hospedagem, sem confundir endereço da plataforma."""
    lodging_types = {'Hotel', 'LodgingBusiness', 'Hostel', 'Motel', 'Resort',
                     'BedAndBreakfast', 'VacationRental', 'Accommodation'}
    if isinstance(value, list):
        for child in value:
            address = address_from_jsonld(child)
            if address:
                return address
    if isinstance(value, dict):
        types = value.get('@type', [])
        if isinstance(types, str):
            types = [types]
        if lodging_types.intersection(types):
            address = value.get('address')
            if isinstance(address, str):
                return address.strip()
            if isinstance(address, dict):
                parts = []
                def normalizar(text):
                    import unicodedata
                    text = unicodedata.normalize('NFKD', str(text).casefold())
                    return re.sub(r'[^a-z0-9]', '', text)
                for key in ('streetAddress', 'addressLocality', 'addressRegion', 'postalCode', 'addressCountry'):
                    part = address.get(key)
                    if isinstance(part, dict):
                        part = part.get('name')
                    if not part:
                        continue
                    part = str(part).strip()
                    normalized = normalizar(part)
                    if normalized and not any(normalized in normalizar(old) for old in parts):
                        parts = [old for old in parts if normalizar(old) not in normalized]
                        parts.append(part)
                if parts:
                    return ', '.join(parts)
        for key in ('@graph', 'mainEntity', 'mainEntityOfPage'):
            address = address_from_jsonld(value.get(key))
            if address:
                return address
    return ""


def fetch_details(driver, url):
    import json
    driver.get(url)
    latitude = longitude = None
    try:
        element = WebDriverWait(driver, 10).until(
            EC.presence_of_element_located((By.ID, "map_trigger_header_pin")))
        latlng = element.get_attribute("data-atlas-latlng")
        if latlng:
            latitude, longitude = map(float, latlng.split(','))
            if not (-90 <= latitude <= 90 and -180 <= longitude <= 180):
                latitude = longitude = None
    except Exception:
        pass

    address = ""
    for element in driver.find_elements(By.CSS_SELECTOR, 'script[type="application/ld+json"]'):
        try:
            address = address_from_jsonld(json.loads(element.get_attribute('textContent') or 'null'))
        except (ValueError, TypeError):
            continue
        if address:
            break
    if not address:
        address = first_text(driver, ['[data-testid="address"]',
                                     '[data-testid="property-address"]',
                                     '#hotel_address', '.hp_address_subtitle',
                                     '[data-node_tt_id="location_score_tooltip"]'])
    return latitude, longitude, address or NOT_INFORMED


def calculate_distance(hotel_coords, event_coords):
    from geopy.distance import geodesic
    try:
        return round(geodesic(hotel_coords, event_coords).kilometers, 2)
    except (ValueError, TypeError):
        return None


def validate_search(destination, checkin, checkout):
    from datetime import datetime
    if not destination.strip():
        raise ValueError('Informe o destino.')
    for value in (checkin, checkout):
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
            raise ValueError('Use datas no formato YYYY-MM-DD.')
    arrival = datetime.strptime(checkin, '%Y-%m-%d')
    departure = datetime.strptime(checkout, '%Y-%m-%d')
    if departure <= arrival:
        raise ValueError('O check-out deve ser posterior ao check-in.')
    return (departure - arrival).days


def is_petropolis(destination):
    import unicodedata
    value = ''.join(c for c in unicodedata.normalize('NFD', destination.lower().strip())
                    if unicodedata.category(c) != 'Mn')
    return value == 'petropolis'


def query_matches(current_url, params):
    from urllib.parse import urlparse, parse_qs
    import unicodedata
    parsed = urlparse(current_url)
    if parsed.scheme != 'https' or parsed.hostname not in ('www.booking.com', 'booking.com'):
        return False
    values = parse_qs(parsed.query)
    def normalized(text):
        return ' '.join(''.join(c for c in unicodedata.normalize('NFD', text.casefold())
                               if unicodedata.category(c) != 'Mn').split())
    for key in ('ss', 'checkin', 'checkout', 'group_adults', 'group_children', 'no_rooms'):
        found = values.get(key, [])
        if len(found) != 1:
            return False
        if key == 'ss':
            if normalized(found[0]) != normalized(str(params[key])):
                return False
        elif found[0] != str(params[key]):
            return False
    return True


def fill_search_form(driver, params):
    from selenium.webdriver.common.keys import Keys
    from datetime import date
    from selenium.common.exceptions import TimeoutException
    wait = WebDriverWait(driver, 15)
    def visible(selector):
        return next((e for e in driver.find_elements(By.CSS_SELECTOR, selector)
                     if e.is_displayed() and e.is_enabled()), False)
    def calendar():
        button = wait.until(lambda d: visible('[data-testid="searchbox-dates-container"]'))
        if button.get_attribute('aria-expanded') != 'true':
            button.click()
        wait.until(lambda d: visible('[data-date]'))
    calendar()
    for key in ('checkin', 'checkout'):
        target = date.fromisoformat(params[key]).isoformat()
        for step in range(25):
            day = visible('[data-date="' + target + '"]')
            if day:
                if day.get_attribute('aria-disabled') == 'true':
                    raise RuntimeError('Data indisponível no calendário: ' + target)
                day.click()
                print('Data selecionada no formulário: ' + target, flush=True)
                break
            shown = [e.get_attribute('data-date') for e in driver.find_elements(By.CSS_SELECTOR, '[data-date]') if e.is_displayed()]
            if not shown or target < min(shown) or step == 24:
                raise RuntimeError('Data fora do calendário acessível: ' + target)
            button = visible('button[aria-label="Mês seguinte"]')
            if not button:
                raise RuntimeError('Botão de próximo mês não encontrado.')
            button.click()
            wait.until(lambda d: [e.get_attribute('data-date') for e in d.find_elements(By.CSS_SELECTOR, '[data-date]') if e.is_displayed()] != shown)
    # Reabre para conferir os dois extremos selecionados antes de enviar.
    calendar()
    for key in ('checkin', 'checkout'):
        day = visible('[data-date="' + params[key] + '"]')
        if not day or day.get_attribute('aria-checked') != 'true':
            raise RuntimeError('Não foi possível conferir as datas selecionadas no formulário.')
    visible('[data-testid="searchbox-dates-container"]').click()
    # O diagnóstico mostrou um campo vazio após digitação; conferir o valor estável.
    retained = False
    for attempt in range(2):
        field = wait.until(lambda d: visible('input[name="ss"]'))
        field.click()
        field.send_keys(Keys.CONTROL, 'a')
        field.send_keys(Keys.BACKSPACE)
        field.send_keys(params['ss'])
        time.sleep(1)
        field = visible('input[name="ss"]')
        if field and field.get_attribute('value').strip().casefold() == params['ss'].strip().casefold():
            retained = True
            break
    if not retained:
        raise RuntimeError('Booking apagou o destino digitado; formulário não enviado.')
    field.send_keys(Keys.ESCAPE)
    field.send_keys(Keys.TAB)
    if visible('input[name="ss"]').get_attribute('value').strip().casefold() != params['ss'].strip().casefold():
        raise RuntimeError('Destino mudou ao sair do campo; formulário não enviado.')
    button = wait.until(lambda d: visible('form[aria-label="Buscar propriedades"] button[type="submit"]'))
    print('Destino preenchido e datas conferidas; enviando pelo formulário...', flush=True)
    button.click()


def open_confirmed_search(driver, url, params):
    from selenium.common.exceptions import TimeoutException
    driver.get(url)
    def ready(d):
        cards = d.find_elements(By.CSS_SELECTOR, '[data-testid="property-card"]')
        return cards if cards and query_matches(d.current_url, params) else False
    try:
        return WebDriverWait(driver, 20).until(ready)
    except TimeoutException as exc:
        if query_matches(driver.current_url, params):
            raise RuntimeError('Consulta manteve os parâmetros, mas não apresentou cartões em 20 segundos.') from exc
    print('Parâmetros perdidos na URL; recuperando a consulta pelo formulário...', flush=True)
    fill_search_form(driver, params)
    try:
        cards = WebDriverWait(driver, 20).until(ready)
    except TimeoutException as exc:
        raise RuntimeError('Busca pelo formulário não retornou cartões com destino, datas e ocupação correspondentes. Nenhum preço foi aceito.') from exc
    print('Consulta do formulário conferida na URL; cartões disponíveis.', flush=True)
    return cards


def write_hotels_excel(path, rows, engine=None):
    """Grava os registros diretamente, sem montar um DataFrame."""
    columns = list(dict.fromkeys(key for row in rows for key in row))
    if engine is None:
        from importlib.util import find_spec
        engine = 'xlsxwriter' if find_spec('xlsxwriter') else 'openpyxl'
    if engine == 'xlsxwriter':
        import xlsxwriter
        with xlsxwriter.Workbook(str(path), {'constant_memory': True, 'strings_to_formulas': False,
                                             'strings_to_urls': False}) as workbook:
            sheet = workbook.add_worksheet('Sheet1')
            header = workbook.add_format({'bold': True})
            for col, name in enumerate(columns):sheet.write_string(0, col, name, header)
            for number, row in enumerate(rows, 1):
                for col, name in enumerate(columns):
                    value = row.get(name)
                    if value is None:continue
                    if isinstance(value, str):sheet.write_string(number, col, value)
                    else:sheet.write(number, col, value)
    elif engine == 'openpyxl':
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet('Sheet1')
        def cell(value, bold=False):
            result = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):result.data_type = 's'
            if bold:result.font = Font(bold=True)
            return result
        sheet.append([cell(name, True) for name in columns])
        for row in rows:sheet.append([cell(row.get(name)) for name in columns])
        workbook.save(str(path))
        workbook.close()
    else:
        raise ValueError('Exportador Excel desconhecido.')
    return engine


def run_scraping(destination, checkin, checkout, *, return_records=False, fast=False, headless=False, on_progress=None):
    from datetime import datetime
    from pathlib import Path
    from urllib.parse import urlencode
    import os
    import tempfile

    destination = destination.strip()
    nights = validate_search(destination, checkin, checkout)
    reference = REFERENCE_COORDS if is_petropolis(destination) else None
    params = dict(ss=destination, checkin=checkin, checkout=checkout,
                  group_adults=2, no_rooms=1, group_children=0,
                  nflt='ht_id=204', selected_currency='BRL')
    url = 'https://www.booking.com/searchresults.pt-br.html?' + urlencode(params)
    began = time.monotonic()
    timings = {}
    print("Abrindo Booking em segundo plano." if headless else "Abrindo Booking para diagnóstico visual.", flush=True)
    driver = create_driver(headless=headless)
    timings['driver_seconds'] = round(time.monotonic() - began, 2)
    hotels = []
    try:
        driver.set_page_load_timeout(60)
        navigation_started = time.monotonic()
        open_confirmed_search(driver, url, params)
        timings['navigation_seconds'] = round(time.monotonic() - navigation_started, 2)
        timings["open_seconds"] = round(time.monotonic() - began, 2)
        phase_start = time.monotonic()
        print("Carregando preços dos cartões...", flush=True)
        def preview(rows):
            if not on_progress:return
            stamp = datetime.now().astimezone().isoformat(timespec='seconds')
            prepared = [{**r, 'Destination': destination, 'Check-in': checkin, 'Check-out': checkout,
                         'Nights': nights, 'Adults': 2, 'Children': 0, 'Rooms': 1,
                         'Address': NOT_INFORMED, 'Latitude': None, 'Longitude': None,
                         'Details Status': 'Prévia conferida; busca em andamento', 'Collected At': stamp}
                        for r in rows]
            if rows and 'first_batch_seconds' not in timings:
                timings['first_batch_seconds'] = round(time.monotonic() - began, 2)
            on_progress(prepared)
        if fast:
            hotels = extract_hotels(driver, fast=True, on_progress=preview,
                                    confirm=lambda current_url: query_matches(current_url, params))
        else:
            hotels = extract_hotels(driver, fast=False)
        if not query_matches(driver.current_url, params):
            raise RuntimeError("Parâmetros da consulta mudaram durante a coleta; resultados descartados.")
        timings["cards_seconds"] = round(time.monotonic() - phase_start, 2)
        phase_start = time.monotonic()
        if not hotels:
            raise RuntimeError('Nenhum cartão de hospedagem foi extraído.')
        print(f'\n{len(hotels)} hospedagens encontradas em {destination}.', flush=True)
        for idx, hotel in enumerate(hotels, 1):
            print(f"Processando {idx}/{len(hotels)}: {hotel['Hotel Name']}", flush=True)
            lat = lon = None
            address = NOT_INFORMED
            error = ''
            if not fast and hotel['Hotel URL'] != NOT_INFORMED:
                try:
                    lat, lon, address = fetch_details(driver, hotel['Hotel URL'])
                except Exception as exc:
                    error = str(exc)
                    print(f'  Falha nos detalhes; preço do cartão preservado: {exc}', flush=True)
            if fast:
                detail_status = "Detalhes não consultados: busca de preços"
            elif error:
                detail_status = 'Falha ao consultar detalhes'
            elif address == NOT_INFORMED or lat is None or lon is None:
                detail_status = 'Detalhes parciais'
            else:
                detail_status = 'Endereço e coordenadas extraídos'
            hotel.update({
                'Destination': destination, 'Check-in': checkin, 'Check-out': checkout,
                'Nights': nights, 'Adults': 2, 'Children': 0, 'Rooms': 1,
                'Requested Currency': 'BRL',
                'Price Basis': 'Valor exibido no cartão; conferir Stay Text e Taxes Text',
                'Latitude': lat, 'Longitude': lon, 'Address': address,
                'Geodesic Distance to Reference (Km)': (
                    calculate_distance((lat, lon), reference)
                    if reference and lat is not None and lon is not None else None),
                'Distance Reference': REFERENCE_NAME if reference else 'Não configurada para este destino',
                'Distance Method': 'Geográfica (geodésica), não trajeto por ruas',
                'Details Status': detail_status, 'Details Error': error,
                'Collected At': datetime.now().astimezone().isoformat(timespec='seconds'),
            })
        timings["details_seconds"] = round(time.monotonic() - phase_start, 2)
    except Exception as exc:
        try:
            import json
            import zipfile
            diagnostic_dir = Path.cwd() / 'data' / 'diagnostics'
            diagnostic_dir.mkdir(parents=True, exist_ok=True)
            diagnostic = diagnostic_dir / ('booking-' + datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '.zip')
            metadata = {'error_type': type(exc).__name__, 'message': str(exc), 'requested_url': url}
            with zipfile.ZipFile(diagnostic, 'x', zipfile.ZIP_DEFLATED) as archive:
                for name, getter in [('url', lambda: driver.current_url), ('title', lambda: driver.title)]:
                    try: metadata[name] = getter()
                    except Exception: pass
                try: archive.writestr('page.html', driver.page_source)
                except Exception: pass
                try: archive.writestr('page.png', driver.get_screenshot_as_png())
                except Exception: pass
                archive.writestr('report.json', json.dumps(metadata, ensure_ascii=False, indent=2))
            print('Diagnóstico salvo: ' + str(diagnostic), flush=True)
        except Exception:
            pass
        if headless:
            print("Consulta oculta não concluída. Verifique o erro; o navegador não será aberto automaticamente.", flush=True)
        raise
    finally:
        try:driver.quit()
        except Exception as cleanup_error:print('Navegador já encerrado: '+type(cleanup_error).__name__, flush=True)

    export_started = time.monotonic()

    # Ordena por total com taxas; totais desconhecidos permanecem no fim.
    for hotel in hotels:
        additional, total, status = price_with_taxes(
            hotel['Price'], hotel['Currency'], hotel['Taxes Text'])
        hotel['Additional Taxes (BRL)'] = additional
        hotel['Total Displayed (BRL)'] = total
        hotel['Total Status'] = status
    hotels.sort(key=lambda row: (row['Total Displayed (BRL)'] is None,
                row['Total Displayed (BRL)'] if row['Total Displayed (BRL)'] is not None else 0,
                row['Price'] is None, row['Price'] if row['Price'] is not None else 0))
    safe_name = re.sub(r'[^\w .-]', '_', destination).strip(' .') or 'destino'
    stamp = datetime.now().strftime('%Y%m%d-%H%M%S-%f')
    filename = Path(f'Hotels - {safe_name} - {checkin} - {checkout} - {stamp}.xlsx').resolve()
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=filename.parent, suffix='.xlsx', delete=False) as f:
            temporary = f.name
        export_engine = write_hotels_excel(temporary, hotels)
        os.replace(temporary, filename)
    finally:
        if temporary and os.path.exists(temporary):
            os.unlink(temporary)
    valid = sum(row['Price'] is not None for row in hotels)
    addresses = sum(row.get('Address') not in (None, '', NOT_INFORMED) for row in hotels)
    stars = sum(type(row.get('Stars')) is int for row in hotels)
    totals = sum(row['Total Displayed (BRL)'] is not None for row in hotels)
    print(f'\nExportadas: {len(hotels)} | Com preço: {valid} | Totais: {totals} | Endereços: {addresses} | Estrelas: {stars}', flush=True)
    print(f'Arquivo salvo: {filename}', flush=True)
    print('Exportador Excel: ' + str(export_engine), flush=True)
    timings["export_seconds"] = round(time.monotonic() - export_started, 2)
    timings["total_seconds"] = round(time.monotonic() - began, 2)
    print("Tempos por etapa: " + str(timings), flush=True)
    if return_records:
        return {"records": hotels, "excel": str(filename), "search_url": url, "fast": fast, "headless": headless, "timings": timings}
    return len(hotels), str(filename)


def on_submit():
    destination = entry_destination.get()
    checkin = entry_checkin.get()
    checkout = entry_checkout.get()

    if not all([destination, checkin, checkout]):
        messagebox.showerror("Erro", "Preencha todos os campos.")
        return

    try:
        count, file = run_scraping(destination, checkin, checkout)
        messagebox.showinfo("Concluído", f"{count} hospedagens exportadas.\nArquivo salvo: {file}")
    except Exception as e:
        messagebox.showerror("Erro", str(e))

# ==================== INTERFACE ====================
if __name__ == "__main__":
    root = tk.Tk()
    root.title("Booking Hotel Scraper")
    root.geometry("520x460")          # ← aumentei a altura
    root.configure(bg="#1e1e2e")
    root.resizable(False, False)

    # Cores
    BG = "#1e1e2e"
    FG = "#cdd6f4"
    ACCENT = "#89b4fa"
    ENTRY_BG = "#313244"
    BUTTON_BG = "#89b4fa"
    BUTTON_FG = "#1e1e2e"

    # Fonte
    FONT_TITLE = ("JetBrains Mono", 14, "bold")
    FONT_LABEL = ("JetBrains Mono", 11)
    FONT_ENTRY = ("JetBrains Mono", 12)

    # Título
    title = tk.Label(
        root,
        text="🏨  Booking Hotel Scraper",
        font=FONT_TITLE,
        bg=BG,
        fg=ACCENT
    )
    title.pack(pady=(25, 15))

    # Frame principal
    frame = tk.Frame(root, bg=BG)
    frame.pack(padx=40, fill="x")

    def create_field(parent, label_text):
        lbl = tk.Label(parent, text=label_text, font=FONT_LABEL, bg=BG, fg=FG, anchor="w")
        lbl.pack(fill="x", pady=(10, 3))

        entry = tk.Entry(
            parent,
            font=FONT_ENTRY,
            bg=ENTRY_BG,
            fg=FG,
            insertbackground=FG,
            relief="flat",
            highlightthickness=1,
            highlightbackground="#45475a",
            highlightcolor=ACCENT
        )
        entry.pack(fill="x", ipady=7)
        return entry

    entry_destination = create_field(frame, "Destination")
    entry_checkin = create_field(frame, "Check-in  (YYYY-MM-DD)")
    entry_checkout = create_field(frame, "Check-out (YYYY-MM-DD)")

    # Botão
    btn = tk.Button(
        root,
        text="▶  Start Scraping",
        font=("JetBrains Mono", 12, "bold"),
        bg=BUTTON_BG,
        fg=BUTTON_FG,
        activebackground="#74c7ec",
        activeforeground=BUTTON_FG,
        relief="flat",
        cursor="hand2",
        command=on_submit,
        padx=25,
        pady=12
    )
    btn.pack(pady=(25, 15))

    # Rodapé
    footer = tk.Label(
        root,
        text="Arch Linux  •  Selenium + Booking.com",
        font=("JetBrains Mono", 9),
        bg=BG,
        fg="#6c7086"
    )
    footer.pack(side="bottom", pady=12)

    root.mainloop()
